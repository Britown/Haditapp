import io
import unittest
from datetime import datetime
import openpyxl
from processor_v3 import extract_all_text, reconcile
from variables_processor import process_unmatched_to_df


class ExcelStatementTests(unittest.TestCase):
    def parse(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.cell(8, 2, '28 sep 2026 - 1 sep 2026')
        sheet.cell(12, 2, '$999.999')
        for col, label in enumerate(['Fecha', 'Categoría', 'Descripción', 'Monto'], 2):
            sheet.cell(18, col, label)
        for row_index, row in enumerate(rows, 19):
            for col, value in enumerate(row, 2):
                sheet.cell(row_index, col, value)
        sheet.cell(90, 2, 'Pie de página 2022')
        stream = io.BytesIO()
        stream.name = 'provisoria.xlsx'
        workbook.save(stream)
        return reconcile(extract_all_text([stream], ''), year=2026, month=9)

    def test_explicit_columns_override_description(self):
        rows = self.parse([
            ['28 sep 2026', 'Cargos', 'Cargo por Compra en TIENDA El 25/09/2026 a las 20:41:34., Monto 9.610', '$9.610'],
            ['14 sep 2026', 'Cargos', 'Cargo por Comisión Checkcard. Neto $6.139, IVA $1.166', '$7.305'],
            ['21 sep 2026', 'Cargos', 'Cargo por abono/pago tarjeta por US$44,53 al tipo de cambio $972,25', '$43.294'],
            ['8 sep 2026', 'Abonos', 'Devolución de cargo por compra con fecha 24-08-2026', '$1.190'],
            [datetime(2026, 9, 1), 'Cargos', 'Transferencia el 2026-08-31 a las 15:07', 35000],
        ])
        self.assertEqual(len(rows), 5)
        self.assertEqual([r['Monto'] for r in rows], [9610, 7305, 43294, 1190, 35000])
        self.assertEqual([r['Fecha'] for r in rows], ['28/09/2026', '14/09/2026', '21/09/2026', '08/09/2026', '01/09/2026'])
        self.assertTrue(all(r['Moneda'] == 'CLP' and not r['Estado'] for r in rows))
        self.assertEqual(rows[3]['Tipo'], 'Ingreso')
        self.assertNotEqual(rows[2]['Tipo'], 'Ingreso')
        self.assertEqual(len(process_unmatched_to_df(rows)), 5)

    def test_missing_amount_and_duplicate_movements_are_preserved(self):
        row = ['1 sep 2026', 'Cargos', 'COMERCIO NUEVO', '$1.000']
        rows = self.parse([row, row, ['1 sep 2026', 'Cargos', 'Monto ilegible', None]])
        self.assertEqual(len(rows), 3)
        self.assertNotEqual(rows[0]['id'], rows[1]['id'])
        self.assertIsNone(rows[2]['Monto'])
        self.assertEqual(rows[2]['Tipo'], 'Revisar')
